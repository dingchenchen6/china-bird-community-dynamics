#!/usr/bin/env python3

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def set_run_font(run, family, size=None):
    run.font.name = family
    if size is not None:
        run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), family)


def iter_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def style_docx(path):
    doc = Document(path)
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.35)
        section.right_margin = Cm(2.35)

    style_specs = {
        "Normal": ("Arial Unicode MS", 10.5),
        "Body Text": ("Arial Unicode MS", 10.5),
        "Title": ("Arial Unicode MS", 20),
        "Heading 1": ("Arial Unicode MS", 15),
        "Heading 2": ("Arial Unicode MS", 12.5),
        "Heading 3": ("Arial Unicode MS", 11),
    }
    for name, (family, size) in style_specs.items():
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        style.font.name = family
        style.font.size = Pt(size)
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attr}"), family)

    for paragraph in iter_paragraphs(doc):
        style_name = paragraph.style.name if paragraph.style else "Normal"
        is_heading = style_name.startswith("Heading") or style_name == "Title"
        family = "Arial Unicode MS"
        for run in paragraph.runs:
            set_run_font(run, family)
        paragraph.paragraph_format.space_after = Pt(4 if not is_heading else 6)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    doc.save(path)


if __name__ == "__main__":
    for arg in sys.argv[1:]:
        style_docx(Path(arg))
