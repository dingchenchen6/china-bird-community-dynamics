#!/usr/bin/env python3
# ============================================================
# 35_qa_figure_package_20261011.py
#
# Scientific question / 科学问题:
#   重绘后的图件是否只改变了呈现而未改变任何数值，且四种交付格式均合格？
#   Did the redraw change presentation only, and are all four deliverable formats sound?
#
# Objective / 分析目标:
#   对 figures_community_submission_20261011/ 的每张图核查：文件齐全；PDF 尺寸、字体嵌入、
#   无中文字符；PNG ≥ 300 dpi；TIFF 450 dpi + LZW；PPTX 可被 python-pptx 打开、幻灯片尺寸
#   等于图件尺寸、可编辑形状与文字段计数；源数据与 Shared 逐字节一致；未改动图件与旧版逐像素
#   一致；LibreOffice 渲染 PPTX 与 PDF 的相似度。另把每个 PDF 渲染为 PNG 供人工目检。
#   Per-figure checks of files, PDF, PNG, TIFF, PPTX, source data, pixel regression, PPTX render.
#
# Input data / 输入数据: figures_community_submission_20261011/{main,supporting,protected_area,source_data};
#   submission_packages_20261010/Shared/; figures_top_journal_20260714_v2/
# Main workflow / 主要流程: 清单 → 逐图检查 → 回归比较 → PPTX 渲染比较 → 写 CSV
# Expected output / 预期输出: qa/QA_figures.csv, qa/pptx_inspection.csv, qa/pixel_regression.csv,
#   qa/pptx_render_similarity.csv, qa/render_pdf/*.png, qa/render_pptx/*.png
# Key assumptions / 关键假设: 需要 poppler（pdfinfo/pdffonts/pdftotext/pdftoppm）、LibreOffice、
#   python-pptx、Pillow、numpy；LibreOffice 用独立配置目录（避免与已打开实例冲突）。
# Main packages / 主要包: python-pptx, Pillow, numpy, subprocess
# Output directory / 输出路径: figures_community_submission_20261011/qa
# ============================================================
"""QA for the 2026-10-11 community-dynamics figure package."""

from __future__ import annotations

import csv
import hashlib
import logging
import re
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Dict, List, Optional

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

logging.basicConfig(level=logging.INFO, format="[35-qa] %(message)s")
logger = logging.getLogger(__name__)
Image.MAX_IMAGE_PIXELS = None

ROOT = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else Path.cwd()
PKG = ROOT / "figures_community_submission_20261011"
QA = PKG / "qa"
SHARED = ROOT / "submission_packages_20261010" / "Shared"
JULY = ROOT / "figures_top_journal_20260714_v2"
CJK = re.compile(r"[　-〿一-鿿＀-￯]")
EMU_PER_MM = 36000

# 旧版文件（用于逐像素回归；None = 有意改动、不做像素比较）/ old files for pixel regression
OLD_PNG: Dict[str, Optional[Path]] = {
    "Figure_4_multidiversity_decoupling": SHARED / "Figures" / "Figure_4_multidiversity_decoupling.png",
    "Figure_5_reorganization_and_winners": SHARED / "Figures" / "Figure_5_reorganization_and_winners.png",
}
for _stem in ["fig05_species_trend_classes_500sp", "fig06_naive_vs_corrected_trends",
              "fig07_global_baselga_turnover_nestedness", "fig08_random_forest_driver_importance",
              "fig09_trait_environment_mechanisms_500sp"]:
    OLD_PNG[f"{_stem}_top_journal_editable"] = JULY / f"{_stem}_top_journal_editable.png"


def run(cmd: List[str]) -> str:
    """Run a command and return stdout (raises on failure)."""
    return subprocess.run(cmd, check=True, capture_output=True, text=True).stdout


def md5(path: Path) -> str:
    """MD5 of a file."""
    return hashlib.md5(path.read_bytes()).hexdigest()


def pdf_checks(pdf: Path) -> Dict[str, object]:
    """Page size (mm), font embedding and CJK text in a PDF."""
    info = run(["pdfinfo", str(pdf)])
    m = re.search(r"Page size:\s+([\d.]+) x ([\d.]+) pts", info)
    w_mm, h_mm = (float(m.group(1)) * 25.4 / 72, float(m.group(2)) * 25.4 / 72) if m else (np.nan, np.nan)
    fonts = run(["pdffonts", str(pdf)]).splitlines()[2:]
    not_embedded = [f.split()[0] for f in fonts if len(f.split()) > 4 and f.split()[-5] != "yes"]
    names = sorted({f.split()[0].split("+")[-1] for f in fonts if f.strip()})
    text = run(["pdftotext", "-q", str(pdf), "-"])
    return {"pdf_w_mm": round(w_mm, 1), "pdf_h_mm": round(h_mm, 1), "pdf_fonts": ";".join(names),
            "pdf_fonts_all_embedded": not not_embedded, "pdf_has_cjk": bool(CJK.search(text))}


def image_checks(png: Path, tiff: Path) -> Dict[str, object]:
    """Resolution and compression of PNG and TIFF."""
    with Image.open(png) as im:
        png_dpi = round(float(im.info.get("dpi", (0, 0))[0]))
        png_px = f"{im.size[0]}x{im.size[1]}"
    with Image.open(tiff) as im:
        tiff_dpi = round(float(im.info.get("dpi", (0, 0))[0]))
        tiff_comp = im.info.get("compression", "")
    return {"png_dpi": png_dpi, "png_px": png_px, "png_ok": png_dpi >= 300,
            "tiff_dpi": tiff_dpi, "tiff_compression": tiff_comp,
            "tiff_ok": tiff_dpi == 450 and tiff_comp == "tiff_lzw"}


def walk(shapes):
    """Yield shapes recursively through groups."""
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes)


def pptx_checks(pptx: Path) -> Dict[str, object]:
    """Load with python-pptx; count editable shapes, text runs and pictures."""
    prs = Presentation(str(pptx))
    slides = list(prs.slides)
    n = pics = runs = 0
    cjk = False
    for sh in walk(slides[0].shapes):
        n += 1
        pics += int(sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if r.text.strip():
                        runs += 1
                        cjk = cjk or bool(CJK.search(r.text))
    return {"pptx_loads": True, "pptx_slides": len(slides),
            "slide_w_mm": round(prs.slide_width / EMU_PER_MM, 1), "slide_h_mm": round(prs.slide_height / EMU_PER_MM, 1),
            "pptx_shapes": n, "pptx_text_runs": runs, "pptx_pictures": pics, "pptx_has_cjk": cjk,
            "pptx_mb": round(pptx.stat().st_size / 1e6, 2)}


def render_pdf(pdf: Path, out_dir: Path, dpi: int = 150) -> Path:
    """Render page 1 of a PDF to PNG for visual inspection."""
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = out_dir / pdf.stem
    run(["pdftoppm", "-r", str(dpi), "-png", "-singlefile", str(pdf), str(stem)])
    return stem.with_suffix(".png")


def pixel_compare(a: Path, b: Path) -> Dict[str, object]:
    """Pixel identity of two PNGs (decoded RGBA)."""
    with Image.open(a) as ia, Image.open(b) as ib:
        x, y = np.asarray(ia.convert("RGBA")), np.asarray(ib.convert("RGBA"))
    if x.shape != y.shape:
        return {"identical": False, "detail": f"shape {x.shape} vs {y.shape}"}
    d = np.abs(x.astype(int) - y.astype(int))
    return {"identical": bool(d.max() == 0), "detail": f"max diff {d.max()}, differing px {(d.sum(2) > 0).mean():.6f}"}


def render_pptx(pptx: Path, out_dir: Path, profile: Path, dpi: int = 150) -> Optional[Path]:
    """Render a PPTX with LibreOffice (isolated profile) to PDF, then PNG."""
    out_dir.mkdir(parents=True, exist_ok=True)
    try:
        run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf",
             "--outdir", str(out_dir), str(pptx)])
    except subprocess.CalledProcessError as err:
        logger.error("LibreOffice failed for %s: %s", pptx.name, err)
        return None
    pdf = out_dir / f"{pptx.stem}.pdf"
    return render_pdf(pdf, out_dir, dpi) if pdf.exists() else None


def similarity(a: Path, b: Path) -> Dict[str, object]:
    """Mean absolute difference (0-255) between two renders after resizing to a common size."""
    with Image.open(a) as ia, Image.open(b) as ib:
        size = (min(ia.size[0], ib.size[0]), min(ia.size[1], ib.size[1]))
        x = np.asarray(ia.convert("L").resize(size), dtype=float)
        y = np.asarray(ib.convert("L").resize(size), dtype=float)
    mad = float(np.abs(x - y).mean())
    return {"mean_abs_diff": round(mad, 2), "share_px_diff_gt64": round(float((np.abs(x - y) > 64).mean()), 4)}


def write_csv(rows: List[Dict[str, object]], path: Path) -> None:
    """Write a list of dicts to CSV."""
    keys: List[str] = []
    for r in rows:
        keys += [k for k in r if k not in keys]
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=keys)
        w.writeheader()
        w.writerows(rows)


def main() -> None:
    """Run all checks and write the QA tables."""
    rows: List[Dict[str, object]] = []
    regress: List[Dict[str, object]] = []
    sims: List[Dict[str, object]] = []
    profile = Path(tempfile.mkdtemp(prefix="lo_profile_"))
    for folder in ["main", "supporting", "protected_area"]:
        for pdf in sorted((PKG / folder).glob("*.pdf")):
            stem = pdf.stem
            files = {ext: PKG / folder / f"{stem}.{ext}" for ext in ["pdf", "png", "tiff", "pptx"]}
            row: Dict[str, object] = {"folder": folder, "figure": stem,
                                      "all_four_formats": all(p.exists() for p in files.values())}
            row.update(pdf_checks(files["pdf"]))
            row.update(image_checks(files["png"], files["tiff"]))
            row.update(pptx_checks(files["pptx"]))
            row["slide_matches_pdf"] = (abs(row["slide_w_mm"] - row["pdf_w_mm"]) < 0.6 and
                                        abs(row["slide_h_mm"] - row["pdf_h_mm"]) < 0.6)
            ref = render_pdf(files["pdf"], QA / "render_pdf")
            out = render_pptx(files["pptx"], QA / "render_pptx", profile)
            if out is not None:
                sims.append({"figure": stem, **similarity(ref, out)})
            if stem in OLD_PNG and OLD_PNG[stem] is not None and OLD_PNG[stem].exists():
                regress.append({"figure": stem, "compared_with": str(OLD_PNG[stem].relative_to(ROOT)),
                                **pixel_compare(files["png"], OLD_PNG[stem])})
            rows.append(row)
            logger.info("%s/%s checked", folder, stem)
    for src in sorted((PKG / "source_data").glob("SourceData_Figure_*.csv")):
        old = SHARED / "Figures" / "source_data" / src.name
        regress.append({"figure": src.stem, "compared_with": str(old.relative_to(ROOT)),
                        "identical": old.exists() and md5(src) == md5(old), "detail": "md5 of CSV"})
    write_csv(rows, QA / "QA_figures.csv")
    write_csv(regress, QA / "pixel_regression.csv")
    write_csv(sims, QA / "pptx_render_similarity.csv")
    write_csv([{k: r[k] for k in ["figure", "slide_w_mm", "slide_h_mm", "pptx_shapes", "pptx_text_runs",
                                  "pptx_pictures", "pptx_mb"]} for r in rows], QA / "pptx_inspection.csv")
    bad = [r["figure"] for r in rows if not (r["all_four_formats"] and r["png_ok"] and r["tiff_ok"]
                                             and r["slide_matches_pdf"] and not r["pdf_has_cjk"]
                                             and not r["pptx_has_cjk"] and r["pdf_fonts_all_embedded"])]
    logger.info("figures checked: %d; failing: %s", len(rows), bad or "none")


if __name__ == "__main__":
    main()
